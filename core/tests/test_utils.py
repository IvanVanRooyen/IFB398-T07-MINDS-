import hashlib
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase

from core.utils import chunk_text, extract_text, sha256_file


class Sha256FileTests(SimpleTestCase):
    def test_hashes_uploaded_file(self):
        content = b"hello world"
        f = SimpleUploadedFile("test.txt", content)
        self.assertEqual(sha256_file(f), hashlib.sha256(content).hexdigest())

    def test_hashes_plain_file_like(self):
        # BytesIO does not contain the `.chunks()` method, which should trigger the
        # fallback
        content = b"some binary content " * 1000
        self.assertEqual(
            sha256_file(io.BytesIO(content)),
            hashlib.sha256(content).hexdigest(),
        )

    def test_restores_file_position(self):
        f = SimpleUploadedFile("test.bin", b"abcdefgh")
        f.seek(3)
        sha256_file(f)
        self.assertEqual(f.tell(), 3)

    def test_empty_file(self):
        f = SimpleUploadedFile("empty.txt", b"")
        self.assertEqual(sha256_file(f), hashlib.sha256(b"").hexdigest())

    def test_large_content(self):
        # multi-chunk hashing tests (`content` should be more than 8KB)
        content = b"x" * (1024 * 1024)
        f = SimpleUploadedFile("big.bin", content)
        self.assertEqual(sha256_file(f), hashlib.sha256(content).hexdigest())


class ChunkTextTests(SimpleTestCase):
    def test_empty_string(self):
        self.assertEqual(chunk_text(""), [])

    def test_whitespace_only(self):
        self.assertEqual(chunk_text("   \n\t  "), [])

    def test_short_text_single_chunk(self):
        self.assertEqual(chunk_text("hello world"), ["hello world"])

    def test_overlap_between_chunks(self):
        words = [f"w{i}" for i in range(120)]
        chunks = chunk_text(" ".join(words), chunk_size=50, overlap=10)
        # step = 50 - 10 = 40, so starts are 0, 40, 80
        self.assertEqual(len(chunks), 3)
        self.assertEqual(chunks[0].split()[0], "w0")
        self.assertEqual(chunks[0].split()[-1], "w49")
        self.assertEqual(chunks[1].split()[0], "w40")
        self.assertEqual(chunks[2].split()[0], "w80")

    def test_zero_overlap(self):
        words = [f"w{i}" for i in range(100)]
        chunks = chunk_text(" ".join(words), chunk_size=25, overlap=0)
        self.assertEqual(len(chunks), 4)
        self.assertEqual(chunks[1].split()[0], "w25")


class ExtractTextTests(SimpleTestCase):
    def test_unsupported_extension_returns_empty(self):
        f = SimpleUploadedFile("notes.txt", b"hello")
        self.assertEqual(extract_text(f), "")

    def test_corrupt_pdf_returns_empty(self):
        # Function swallows exceptions and returns ""
        f = SimpleUploadedFile("bad.pdf", b"not a real pdf")
        self.assertEqual(extract_text(f), "")

    def test_pdf_extraction(self):
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        c.drawString(100, 750, "Hello PDF World")
        c.showPage()
        c.save()

        f = SimpleUploadedFile("sample.pdf", buf.getvalue())
        self.assertIn("Hello PDF World", extract_text(f))

    def test_docx_extraction(self):
        from docx import Document

        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        doc.save(buf)

        f = SimpleUploadedFile("sample.docx", buf.getvalue())
        text = extract_text(f)
        self.assertIn("First paragraph", text)
        self.assertIn("Second paragraph", text)

    def test_file_pointer_reset_after_extraction(self):
        # The comment in extract_text says "reset so the file saves correctly"
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        c.drawString(100, 750, "x")
        c.showPage()
        c.save()

        f = SimpleUploadedFile("sample.pdf", buf.getvalue())
        extract_text(f)
        self.assertEqual(f.tell(), 0)
